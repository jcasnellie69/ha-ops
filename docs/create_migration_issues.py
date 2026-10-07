#!/usr/bin/env python3
# D100726T0227 | HA-0001 | correct migration script paths | JC | ha-ops repo
"""
z/OS 3.2 Migration — GitHub Issue & Project Board Automation
=============================================================
Creates GitHub Issues from the z/OS migration Excel workbook and
optionally adds them to a GitHub Projects v2 board with sprint
assignments, labels, and full migration detail.

Usage:
  # Dry run (preview only, no API calls)
  python docs/create_migration_issues.py --token $GH_TOKEN --repo owner/repo --dry-run

  # PARMLIB sheet only
  python docs/create_migration_issues.py --token $GH_TOKEN --repo owner/repo --sheets "PARMLIB Changes"

  # All sheets
  python docs/create_migration_issues.py --token $GH_TOKEN --repo owner/repo

  # All sheets + add to Project board
  python docs/create_migration_issues.py --token $GH_TOKEN --repo owner/repo --project-id PVT_xxx

Requirements:
  pip install openpyxl requests python-dotenv
"""

import argparse
import json
import os
import sys
import time
from pathlib import Path

try:
    import openpyxl
    import requests
except ImportError:
    print("❌ Missing dependencies. Run: pip install openpyxl requests")
    sys.exit(1)

# ── Constants ─────────────────────────────────────────────────────────────────

DEFAULT_WORKBOOK = str(Path(__file__).resolve().with_name("zos-migration-risk-comparison.xlsx"))

GITHUB_API      = "https://api.github.com"
GITHUB_GRAPHQL  = "https://api.github.com/graphql"

# Sprint assignment: priority/risk → sprint label
SPRINT_BY_RISK = {
    "ipl failure": "sprint-1",
    "critical":    "sprint-1",
    "high":        "sprint-2",
    "important":   "sprint-3",
    "medium":      "sprint-3",
    "beneficial":  "sprint-5",
    "low":         "sprint-5",
}

SPRINT_BY_FOCUS = {
    "security": "sprint-2",
    "qsc":      "sprint-2",
    "ai":       "sprint-4",
    "cloud":    "sprint-5",
    "monitor":  "sprint-5",
    "data":     "sprint-5",
    "ops":      "sprint-3",
}

# Label definitions: name → (color_hex, description)
LABEL_DEFINITIONS = {
    "parmlib":         ("0075ca", "PARMLIB member change required"),
    "ipl-blocker":     ("d73a4a", "Causes IPL failure if not addressed before IPL"),
    "security":        ("e4e669", "Security / RACF / ICSF change"),
    "ai-workload":     ("a2eeef", "AI workload enablement"),
    "deprecated":      ("cfd3d7", "Deprecated or removed function"),
    "qsc":             ("7057ff", "Quantum-safe cryptography"),
    "hybrid-cloud":    ("0e8a16", "Hybrid cloud integration"),
    "monitoring":      ("1d76db", "Monitoring / RMF / SMF / CDP"),
    "critical":        ("b60205", "Critical priority — must complete before IPL"),
    "important":       ("e99695", "Important priority — complete before production"),
    "beneficial":      ("c2e0c6", "Beneficial / enhancement"),
    "sprint-1":        ("f9d0c4", "Sprint 1: IPL Blockers & Foundation"),
    "sprint-2":        ("fef2c0", "Sprint 2: Security Architecture"),
    "sprint-3":        ("d4edda", "Sprint 3: PARMLIB Hardening"),
    "sprint-4":        ("cce5ff", "Sprint 4: AI Workload Enablement"),
    "sprint-5":        ("e2d9f3", "Sprint 5: Cloud & Observability"),
    "sprint-6":        ("f8d7da", "Sprint 6: Validation & Cutover"),
    "needs-review":    ("fbca04", "Requires system programmer review"),
    "migration-log":   ("006b75", "Must be documented in migration log"),
}

# ── Label mapping helpers ─────────────────────────────────────────────────────

def resolve_labels(priority: str, focus: str, status: str,
                   component: str, ipl_risk: str) -> list[str]:
    labels = set()
    p = (priority or "").lower()
    f = (focus    or "").lower()
    s = (status   or "").lower()
    c = (component or "").lower()
    r = (ipl_risk  or "").lower()

    # Priority
    if "ipl failure" in p or "yes" in r:
        labels.update(["ipl-blocker", "critical", "parmlib"])
    elif "critical" in p:
        labels.add("critical")
    elif "high" in p or "important" in p:
        labels.add("important")
    elif "beneficial" in p or "low" in p:
        labels.add("beneficial")

    # Focus area
    if "security" in f or "qsc" in f:
        labels.add("security")
    if "qsc" in f:
        labels.add("qsc")
    if "ai" in f:
        labels.add("ai-workload")
    if "cloud" in f:
        labels.add("hybrid-cloud")
    if "monitor" in f:
        labels.add("monitoring")

    # Status
    if "removed" in s or "deprecated" in s:
        labels.add("deprecated")
    if "changed default" in s or "parmlib" in s:
        labels.add("parmlib")

    # Component heuristic
    if any(x in c for x in ["xx", "parmlib", "icsf", "csf", "jes", "bpx",
                              "iea", "irr", "smf", "diag", "iar", "izu"]):
        labels.add("parmlib")

    labels.add("migration-log")
    return sorted(labels)


def resolve_sprint(priority: str, focus: str, ipl_risk: str) -> str:
    if (ipl_risk or "").upper() == "YES":
        return "sprint-1"
    p = (priority or "").lower()
    f = (focus    or "").lower()
    for key, sprint in SPRINT_BY_FOCUS.items():
        if key in f:
            return sprint
    for key, sprint in SPRINT_BY_RISK.items():
        if key in p:
            return sprint
    return "sprint-3"


# ── Issue body builder ────────────────────────────────────────────────────────

def build_issue_body(data: dict, sheet: str) -> str:
    lines = [f"## z/OS 3.2 Migration Task — {sheet}\n"]

    def row(icon, label, val):
        v = str(val or "").strip()
        if v and v not in ("-", "None", "☐", "nan"):
            lines.append(f"| {icon} **{label}** | {v} |")

    lines.append("| Field | Value |")
    lines.append("|---|---|")
    row("📄", "Component / Member",  data.get("Component") or data.get("Member"))
    row("🗓️", "Release / Since",     data.get("Release") or data.get("Since"))
    row("🚨", "Priority / Impact",   data.get("Priority") or data.get("Impact"))
    row("📌", "Status",              data.get("Status"))
    row("💥", "IPL Failure Risk",    data.get("IPL Failure Risk"))
    row("🎯", "Focus Area",          data.get("Focus Area"))
    row("🤖", "AI Impact",           data.get("AI Impact"))
    row("🔐", "QSC Impact",          data.get("QSC Impact"))
    row("☁️", "Cloud Impact",        data.get("Cloud Impact"))
    row("⏱️", "Effort Estimate",     data.get("Effort Estimate"))

    detail = data.get("Key Detail") or data.get("Risk Detail") or data.get("Detail")
    if detail and str(detail).strip() not in ("None", "-"):
        lines.append(f"\n### 📋 Detail\n{detail}")

    action = data.get("Required Action") or data.get("Replacement / Action")
    if action and str(action).strip() not in ("None", "-"):
        lines.append(f"\n### ✅ Required Action\n```\n{action}\n```")

    lines.append("""
### Acceptance Criteria
- [ ] Change implemented and tested in **non-production LPAR**
- [ ] Reviewed by system programmer
- [ ] Documented in migration log (`docs/migration-log.md`)
- [ ] Regression tested — no new abends or errors
- [ ] Signed off for production IPL

### Notes
<!-- Add implementation notes, test results, or blockers here -->

---
*Auto-generated from z/OS 3.2 Migration Workbook — do not edit the table above manually*""")

    return "\n".join(lines)


# ── Sheet parsers ─────────────────────────────────────────────────────────────

def parse_parmlib_sheet(ws) -> list[dict]:
    issues = []
    headers = [str(c.value or "").strip() for c in ws[4]]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        d = dict(zip(headers, row))
        member  = str(d.get("Member", "") or "").strip()
        change  = str(d.get("Parameter / Change", "") or "").strip()
        risk    = str(d.get("Risk Level", "") or "").strip()
        ipl     = str(d.get("IPL Failure Risk", "") or "").strip()
        action  = str(d.get("Required Action", "") or "").strip()

        title = f"[PARMLIB] {member} — {change[:75]}"
        data  = {
            "Component":       member,
            "Member":          member,
            "Release":         d.get("Release", ""),
            "Priority":        risk,
            "Status":          d.get("Change Type", ""),
            "IPL Failure Risk": ipl,
            "AI Impact":       d.get("AI Workload Impact", ""),
            "QSC Impact":      d.get("QSC Impact", ""),
            "Cloud Impact":    d.get("Cloud Impact", ""),
            "Key Detail":      change,
            "Required Action": action,
        }
        labels = resolve_labels(risk, "", d.get("Change Type",""), member, ipl)
        sprint = resolve_sprint(risk, "", ipl)
        issues.append({"title": title, "body": build_issue_body(data, "PARMLIB Changes"),
                        "labels": labels, "sprint": sprint, "sheet": "PARMLIB Changes"})
    return issues


def parse_deprecated_sheet(ws) -> list[dict]:
    issues = []
    headers = [str(c.value or "").strip() for c in ws[4]]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        d       = dict(zip(headers, row))
        feature = str(d.get("Component / Feature", "") or "").strip()
        impact  = str(d.get("Impact Level", "") or "").strip()
        status  = str(d.get("Status", "") or "").strip()
        ipl     = str(d.get("IPL Failure Risk", "") or "").strip()

        title = f"[DEPRECATED] {feature[:90]}"
        data  = {
            "Component":       feature[:40],
            "Since":           d.get("Removed In", ""),
            "Priority":        impact,
            "Impact":          impact,
            "Status":          status,
            "IPL Failure Risk": ipl,
            "Effort Estimate": d.get("Effort Estimate", ""),
            "Risk Detail":     d.get("Risk / Detail", ""),
            "Replacement / Action": d.get("Replacement / Action", ""),
        }
        labels = resolve_labels(impact, "", status, feature, ipl)
        sprint = resolve_sprint(impact, "", ipl)
        issues.append({"title": title, "body": build_issue_body(data, "Deprecated & Removed"),
                        "labels": labels, "sprint": sprint, "sheet": "Deprecated & Removed"})
    return issues


def parse_function_matrix_sheet(ws) -> list[dict]:
    issues = []
    headers = [str(c.value or "").strip() for c in ws[4]]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        d         = dict(zip(headers, row))
        component = str(d.get("Component", "") or "").strip()
        feature   = str(d.get("Function / Feature", "") or "").strip()
        priority  = str(d.get("Priority", "") or "").strip()
        focus     = str(d.get("Focus Area", "") or "").strip()
        status    = str(d.get("Status", "") or "").strip()

        title = f"[{component}] {feature[:80]}"
        data  = {
            "Component":  component,
            "Release":    d.get("Release", ""),
            "Priority":   priority,
            "Status":     status,
            "Focus Area": focus,
            "AI Impact":  d.get("AI Impact", ""),
            "QSC Impact": d.get("QSC Impact", ""),
            "Cloud Impact": d.get("Cloud Impact", ""),
            "Key Detail": d.get("Key Detail", ""),
        }
        labels = resolve_labels(priority, focus, status, component, "")
        sprint = resolve_sprint(priority, focus, "")
        issues.append({"title": title, "body": build_issue_body(data, "Function Matrix"),
                        "labels": labels, "sprint": sprint, "sheet": "Function Matrix"})
    return issues


SHEET_PARSERS = {
    "PARMLIB Changes":    parse_parmlib_sheet,
    "Deprecated & Removed": parse_deprecated_sheet,
    "Function Matrix":    parse_function_matrix_sheet,
}


# ── GitHub REST helpers ───────────────────────────────────────────────────────

class GitHubClient:
    def __init__(self, token: str, repo: str, dry_run: bool = False):
        self.token   = token
        self.repo    = repo
        self.dry_run = dry_run
        self.headers = {
            "Authorization":        f"Bearer {token}",
            "Accept":               "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
        }
        self._existing_labels: set[str] = set()

    def _get(self, path: str) -> requests.Response:
        return requests.get(f"{GITHUB_API}{path}", headers=self.headers, timeout=30)

    def _post(self, path: str, payload: dict) -> requests.Response:
        return requests.post(f"{GITHUB_API}{path}", headers=self.headers,
                             json=payload, timeout=30)

    def _patch(self, path: str, payload: dict) -> requests.Response:
        return requests.patch(f"{GITHUB_API}{path}", headers=self.headers,
                              json=payload, timeout=30)

    def graphql(self, query: str, variables: dict = None) -> dict:
        payload = {"query": query}
        if variables:
            payload["variables"] = variables
        resp = requests.post(GITHUB_GRAPHQL, headers=self.headers,
                             json=payload, timeout=30)
        return resp.json()

    # ── Labels ────────────────────────────────────────────────────────────────

    def fetch_existing_labels(self):
        resp = self._get(f"/repos/{self.repo}/labels?per_page=100")
        if resp.status_code == 200:
            self._existing_labels = {l["name"] for l in resp.json()}
        print(f"  Found {len(self._existing_labels)} existing labels")

    def ensure_label(self, name: str, color: str, description: str):
        if name in self._existing_labels:
            return
        if self.dry_run:
            print(f"  [DRY RUN] Would create label: {name}")
            self._existing_labels.add(name)
            return
        resp = self._post(f"/repos/{self.repo}/labels",
                          {"name": name, "color": color, "description": description})
        if resp.status_code == 201:
            self._existing_labels.add(name)
            print(f"  ✅ Created label: {name}")
        elif resp.status_code == 422:
            self._existing_labels.add(name)  # already exists
        else:
            print(f"  ⚠️  Label '{name}' failed ({resp.status_code})")

    def ensure_all_labels(self):
        print("\n🏷️  Ensuring all labels exist...")
        if not self.dry_run:
            self.fetch_existing_labels()
        for name, (color, desc) in LABEL_DEFINITIONS.items():
            self.ensure_label(name, color, desc)

    # ── Milestones ────────────────────────────────────────────────────────────

    def ensure_milestones(self) -> dict[str, int]:
        """Create sprint milestones and return name→number map."""
        milestones_def = [
            ("Sprint 1 — IPL Blockers & Foundation",    "Remove all IPL-blocking deprecated params; establish migration baseline"),
            ("Sprint 2 — Security Architecture",         "ICSF PQC enablement; System SSL audit; RACF enhancements"),
            ("Sprint 3 — PARMLIB Hardening",             "CPENABLE=SYSTEM testing; IARPRMxx; DIAGxx; SMFLIMxx; IRRPRMxx"),
            ("Sprint 4 — AI Workload Enablement",        "AI System Services 1.2; WLM Policy Advisor; zDNN; BUFSIZE tuning"),
            ("Sprint 5 — Cloud & Observability",         "OpenTelemetry; OAM REST API; Docker networking; z/OSMF automation"),
            ("Sprint 6 — Validation & Cutover",          "Full IPL test; SMF record validation; RMF OpenMetrics; production cutover"),
        ]
        result = {}
        if self.dry_run:
            for i, (title, _) in enumerate(milestones_def, 1):
                result[f"sprint-{i}"] = i
                print(f"  [DRY RUN] Would create milestone: {title}")
            return result

        # Fetch existing
        resp = self._get(f"/repos/{self.repo}/milestones?state=open&per_page=50")
        existing = {}
        if resp.status_code == 200:
            existing = {m["title"]: m["number"] for m in resp.json()}

        for i, (title, desc) in enumerate(milestones_def, 1):
            key = f"sprint-{i}"
            if title in existing:
                result[key] = existing[title]
                print(f"  ♻️  Milestone exists: {title} (#{existing[title]})")
            else:
                r = self._post(f"/repos/{self.repo}/milestones",
                               {"title": title, "description": desc, "state": "open"})
                if r.status_code == 201:
                    num = r.json()["number"]
                    result[key] = num
                    print(f"  ✅ Created milestone: {title} (#{num})")
                else:
                    print(f"  ⚠️  Milestone failed ({r.status_code}): {title}")
        return result

    # ── Issues ────────────────────────────────────────────────────────────────

    def create_issue(self, title: str, body: str, labels: list[str],
                     milestone: int | None = None) -> dict | None:
        if self.dry_run:
            print(f"  [DRY RUN] Issue: {title[:70]}")
            return {"number": 0, "html_url": "dry-run", "node_id": "dry-run"}

        payload: dict = {"title": title, "body": body, "labels": labels}
        if milestone:
            payload["milestone"] = milestone

        resp = self._post(f"/repos/{self.repo}/issues", payload)
        if resp.status_code == 201:
            return resp.json()
        print(f"  ❌ Issue failed ({resp.status_code}): {resp.text[:200]}")
        return None

    # ── Projects v2 (GraphQL) ─────────────────────────────────────────────────

    def add_issue_to_project(self, project_id: str, issue_node_id: str) -> str | None:
        """Add an issue to a Projects v2 board. Returns item node_id."""
        mutation = """
        mutation($projectId: ID!, $contentId: ID!) {
          addProjectV2ItemById(input: {projectId: $projectId, contentId: $contentId}) {
            item { id }
          }
        }"""
        result = self.graphql(mutation, {"projectId": project_id,
                                          "contentId": issue_node_id})
        try:
            return result["data"]["addProjectV2ItemById"]["item"]["id"]
        except (KeyError, TypeError):
            return None

    def set_project_field(self, project_id: str, item_id: str,
                          field_id: str, value: dict):
        """Set a single-select or text field on a project item."""
        mutation = """
        mutation($projectId: ID!, $itemId: ID!, $fieldId: ID!, $value: ProjectV2FieldValue!) {
          updateProjectV2ItemFieldValue(input: {
            projectId: $projectId, itemId: $itemId,
            fieldId: $fieldId, value: $value
          }) { projectV2Item { id } }
        }"""
        self.graphql(mutation, {
            "projectId": project_id,
            "itemId":    item_id,
            "fieldId":   field_id,
            "value":     value,
        })

    def get_project_fields(self, project_id: str) -> dict:
        """Fetch all field IDs and option IDs for a project."""
        query = """
        query($projectId: ID!) {
          node(id: $projectId) {
            ... on ProjectV2 {
              fields(first: 30) {
                nodes {
                  ... on ProjectV2Field { id name }
                  ... on ProjectV2SingleSelectField {
                    id name
                    options { id name }
                  }
                  ... on ProjectV2IterationField {
                    id name
                    configuration { iterations { id title startDate } }
                  }
                }
              }
            }
          }
        }"""
        result = self.graphql(query, {"projectId": project_id})
        fields = {}
        try:
            for node in result["data"]["node"]["fields"]["nodes"]:
                if not node:
                    continue
                name = node.get("name", "")
                fields[name] = node
        except (KeyError, TypeError):
            pass
        return fields


# ── Project board setup ───────────────────────────────────────────────────────

def setup_project_board(client: GitHubClient, project_id: str,
                         issue_results: list[dict]):
    """Add issues to project board and set Status/Sprint fields."""
    print(f"\n📋 Setting up Project board (ID: {project_id})...")

    if client.dry_run:
        print("  [DRY RUN] Would add issues to project board")
        return

    fields = client.get_project_fields(project_id)
    print(f"  Found fields: {list(fields.keys())}")

    # Resolve Status field
    status_field = fields.get("Status")
    status_options = {}
    if status_field and "options" in status_field:
        status_options = {o["name"]: o["id"] for o in status_field["options"]}
        print(f"  Status options: {list(status_options.keys())}")

    # Resolve Sprint/Iteration field
    sprint_field = fields.get("Sprint") or fields.get("Iteration")
    sprint_options = {}
    if sprint_field and "configuration" in sprint_field:
        for it in sprint_field["configuration"].get("iterations", []):
            sprint_options[it["title"]] = it["id"]
        print(f"  Sprint iterations: {list(sprint_options.keys())}")

    added = 0
    for item in issue_results:
        node_id = item.get("node_id")
        sprint  = item.get("sprint", "sprint-3")
        if not node_id or node_id == "dry-run":
            continue

        # Add to project
        project_item_id = client.add_issue_to_project(project_id, node_id)
        if not project_item_id:
            continue

        # Set Status → "Todo"
        if status_field and "Todo" in status_options:
            client.set_project_field(project_id, project_item_id,
                                     status_field["id"],
                                     {"singleSelectOptionId": status_options["Todo"]})

        # Set Sprint iteration
        sprint_label_map = {
            "sprint-1": "Sprint 1",
            "sprint-2": "Sprint 2",
            "sprint-3": "Sprint 3",
            "sprint-4": "Sprint 4",
            "sprint-5": "Sprint 5",
            "sprint-6": "Sprint 6",
        }
        sprint_name = sprint_label_map.get(sprint, "Sprint 3")
        if sprint_field and sprint_name in sprint_options:
            client.set_project_field(project_id, project_item_id,
                                     sprint_field["id"],
                                     {"iterationId": sprint_options[sprint_name]})

        added += 1
        time.sleep(0.3)

    print(f"  ✅ Added {added} issues to project board")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="z/OS 3.2 Migration — GitHub Issue & Project Board Automation",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Dry run — preview everything
  python docs/create_migration_issues.py --token $GH_TOKEN --repo org/repo --dry-run

  # PARMLIB sheet only
  python docs/create_migration_issues.py --token $GH_TOKEN --repo org/repo \\
    --sheets "PARMLIB Changes"

  # All sheets + project board
  python docs/create_migration_issues.py --token $GH_TOKEN --repo org/repo \\
    --project-id PVT_kgDOBxxxxxx

  # From .env file
  python docs/create_migration_issues.py --env-file .env
        """
    )
    parser.add_argument("--token",    help="GitHub Personal Access Token (or set GH_TOKEN env var)")
    parser.add_argument("--repo",     help="GitHub repo: owner/repo-name (or set GH_REPO env var)")
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK,
                        help=f"Path to Excel workbook (default: {DEFAULT_WORKBOOK})")
    parser.add_argument("--sheets",   nargs="+",
                        default=["PARMLIB Changes", "Deprecated & Removed", "Function Matrix"],
                        help="Sheets to process")
    parser.add_argument("--project-id", default=None,
                        help="GitHub Projects v2 node ID (PVT_xxx) to add issues to")
    parser.add_argument("--dry-run",  action="store_true",
                        help="Preview without making any API calls")
    parser.add_argument("--delay",    type=float, default=0.6,
                        help="Seconds between API calls (default: 0.6)")
    parser.add_argument("--env-file", default=".env",
                        help="Path to .env file for token/repo (default: .env)")
    parser.add_argument("--output",   default="migration_issues_created.json",
                        help="Output JSON summary file")
    parser.add_argument("--skip-labels",     action="store_true", help="Skip label creation")
    parser.add_argument("--skip-milestones", action="store_true", help="Skip milestone creation")
    args = parser.parse_args()

    # Load .env if present
    env_path = Path(args.env_file)
    if env_path.exists():
        with open(env_path) as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

    token = args.token or os.environ.get("GH_TOKEN") or os.environ.get("GITHUB_TOKEN")
    repo  = args.repo  or os.environ.get("GH_REPO")  or os.environ.get("GITHUB_REPO")

    if not token:
        print("❌ GitHub token required. Use --token or set GH_TOKEN in environment/.env")
        sys.exit(1)
    if not repo:
        print("❌ GitHub repo required. Use --repo owner/repo or set GH_REPO in environment/.env")
        sys.exit(1)

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        print(f"❌ Workbook not found: {workbook_path}")
        sys.exit(1)

    print(f"""
╔══════════════════════════════════════════════════════════════╗
║     z/OS 3.2 Migration — GitHub Automation Script           ║
╚══════════════════════════════════════════════════════════════╝
  Repo:     {repo}
  Workbook: {workbook_path}
  Sheets:   {', '.join(args.sheets)}
  Project:  {args.project_id or 'None'}
  Dry Run:  {args.dry_run}
""")

    client = GitHubClient(token, repo, dry_run=args.dry_run)

    # ── Step 1: Labels ────────────────────────────────────────────────────────
    if not args.skip_labels:
        client.ensure_all_labels()

    # ── Step 2: Milestones ────────────────────────────────────────────────────
    milestones = {}
    if not args.skip_milestones:
        print("\n🏁 Creating sprint milestones...")
        milestones = client.ensure_milestones()

    # ── Step 3: Parse workbook ────────────────────────────────────────────────
    print(f"\n📂 Loading workbook: {workbook_path}")
    wb = openpyxl.load_workbook(workbook_path, data_only=True)

    all_issues = []
    for sheet_name in args.sheets:
        if sheet_name not in wb.sheetnames:
            print(f"⚠️  Sheet '{sheet_name}' not found — skipping")
            continue
        ws     = wb[sheet_name]
        parser = SHEET_PARSERS.get(sheet_name)
        if not parser:
            print(f"⚠️  No parser for sheet '{sheet_name}' — skipping")
            continue
        rows = parser(ws)
        all_issues.extend(rows)
        print(f"✅ Parsed {len(rows):3d} issues from '{sheet_name}'")

    print(f"\n📋 Total issues to create: {len(all_issues)}")
    if args.dry_run:
        print("🔍 DRY RUN — no GitHub API calls will be made\n")

    # ── Step 4: Create issues ─────────────────────────────────────────────────
    print("\n🚀 Creating GitHub Issues...\n")
    created_issues = []
    failed = 0

    # Group by sprint for display
    sprint_counts: dict[str, int] = {}

    for i, issue in enumerate(all_issues, 1):
        sprint    = issue["sprint"]
        labels    = issue["labels"] + [sprint]
        milestone = milestones.get(sprint)

        print(f"  [{i:03d}/{len(all_issues)}] [{sprint}] {issue['title'][:65]}...")

        result = client.create_issue(
            title=issue["title"],
            body=issue["body"],
            labels=list(dict.fromkeys(labels)),
            milestone=milestone,
        )

        if result:
            sprint_counts[sprint] = sprint_counts.get(sprint, 0) + 1
            created_issues.append({
                "number":   result.get("number"),
                "title":    issue["title"],
                "sprint":   sprint,
                "sheet":    issue["sheet"],
                "labels":   labels,
                "url":      result.get("html_url"),
                "node_id":  result.get("node_id"),
            })
        else:
            failed += 1

        if not args.dry_run:
            time.sleep(args.delay)

    # ── Step 5: Add to Project board ──────────────────────────────────────────
    if args.project_id and created_issues:
        setup_project_board(client, args.project_id, created_issues)

    # ── Step 6: Save summary ──────────────────────────────────────────────────
    summary = {
        "repo":          repo,
        "total_created": len(created_issues),
        "total_failed":  failed,
        "dry_run":       args.dry_run,
        "by_sprint":     sprint_counts,
        "issues":        created_issues,
    }
    with open(args.output, "w") as f:
        json.dump(summary, f, indent=2)

    # ── Report ────────────────────────────────────────────────────────────────
    print(f"""
╔══════════════════════════════════════════════════════════════╗
║                        SUMMARY                              ║
╚══════════════════════════════════════════════════════════════╝
  ✅ Created:  {len(created_issues)} issues
  ❌ Failed:   {failed} issues
  📄 Output:   {args.output}

  Issues by Sprint:""")
    for sprint in sorted(sprint_counts):
        bar = "█" * sprint_counts[sprint]
        print(f"    {sprint}: {sprint_counts[sprint]:3d}  {bar}")

    if not args.dry_run and created_issues:
        print(f"\n  🔗 View issues: https://github.com/{repo}/issues")
        if args.project_id:
            print(f"  📋 View board:  https://github.com/orgs/{repo.split('/')[0]}/projects")


if __name__ == "__main__":
    main()